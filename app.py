import base64
from datetime import datetime, date, timedelta
import io
from io import BytesIO
import os
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import pandas as pd
import qrcode
import requests
import json
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

FILE_PATH = "objednavky_deni_b2b.csv"
SETTINGS_PATH = "nastaveni_deni.json"
LOGO_PATH = "logo.jpg"

# --- BANKOVNÍ ÚDAJE DENI CANDLE ---
BANK_ACCOUNT = "4078292013"
BANK_CODE = "0800"
BANK_IBAN = "CZ5808000000004078292013"

GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "")
GITHUB_REPO = st.secrets.get("GITHUB_REPO", "")
ADMIN_PASSWORD = st.secrets.get("ADMIN_PASSWORD", "deni2026")
VÝCHOZÍ_B2B_PIN = st.secrets.get("B2B_PIN", "partner2026")

EMAIL_SENDER = st.secrets.get("EMAIL_SENDER", "")
EMAIL_PASSWORD = st.secrets.get("EMAIL_PASSWORD", "")
EMAIL_RECEIVER = st.secrets.get("EMAIL_RECEIVER", "")
SMTP_SERVER = st.secrets.get("SMTP_SERVER", "smtp.centrum.cz")
SMTP_PORT = st.secrets.get("SMTP_PORT", 465)

# Kompletní katalog produktů
PRODUKTY_KATALOG = [
    # Vánoční kolekce
    {"nazev": "🎄 Šťastné a veselé – stromeček", "cena_mo": 409.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "✨ Krásné prožití vánočních svátků (velká)", "cena_mo": 409.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "🌟 Kouzelné Vánoce", "cena_mo": 409.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "☕ Pohodové Vánoce", "cena_mo": 299.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "🎁 Veselé vánoce – dáreček", "cena_mo": 299.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "🕯️ Krásné a klidné Vánoce", "cena_mo": 299.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "⭐ Vánoční hvězda – červená", "cena_mo": 299.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "❄️ Krásné Vánoce – vločka", "cena_mo": 399.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "✨ Krásné prožití vánočních svátků", "cena_mo": 399.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "❄️ Veselé Vánoce – vločka", "cena_mo": 399.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "🍪 Perníček", "cena_mo": 399.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},
    {"nazev": "🏠 Perníková chaloupka", "cena_mo": 399.0, "foto": "", "kategorie": "🎄 Vánoční kolekce"},

    # Podzim & Halloween
    {"nazev": "🦔 Ježkův podzim", "cena_mo": 359.0, "foto": "", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "🎃 Halloween", "cena_mo": 359.0, "foto": "", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "🎃 Dýňulka", "cena_mo": 269.0, "foto": "dynulka.jpg", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "☁️ Podzimní obláček", "cena_mo": 359.0, "foto": "oblacek.jpg", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "🎃 Dýňový okamžik", "cena_mo": 359.0, "foto": "", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "🍂 Dýňový podzim", "cena_mo": 359.0, "foto": "", "kategorie": "🎃 Podzim & Halloween"},
    {"nazev": "⛄ Dýňulka sněhulka", "cena_mo": 269.0, "foto": "", "kategorie": "🎃 Podzim & Halloween"},

    # Květinové & Elegantní
    {"nazev": "🌿 Zahrada snů", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌸 Pivoňkové objetí", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌹 Sametový květ – růžový", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "💜 Sametový květ – fialový", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌼 Nebeská kopretina", "cena_mo": 249.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "💜 Nebe na dlani", "cena_mo": 249.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "👑 Královská perla – Perleť", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌸 Královská perla – růžová", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🟡 Královská perla – Zlatá", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🪷 Lotosový květ – růžový", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🪷 Lotosový květ – fialový", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "💜 Pastelová elegance – frézie", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌷 Pastelová elegance – tulipány", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "💜 Jarní pohlazení", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌼 Barevná zahrádka", "cena_mo": 349.0, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "💜 Fialový květ", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌹 Růžový květ", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},
    {"nazev": "🌼 Žlutý květ", "cena_mo": 249.9, "foto": "", "kategorie": "🌸 Květinové & Elegantní"},

    # Dárkové & Motivační
    {"nazev": "🍋 Citronová pomsta (citronela)", "cena_mo": 259.9, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "❤️ Děkuji!", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "🌸 Chvíle pro tebe", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "🌺 Svítím pro tebe", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "🤍 Pro radost…", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "🕯️ Rozsviť si den", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "✨ Jen tak…", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},
    {"nazev": "🌙 Vypni svět, zapal svíčku", "cena_mo": 259.0, "foto": "", "kategorie": "🕯️ Dárkové & Motivační"},

    # Láska & Zvířátka
    {"nazev": "💖 Srdíčko z lásky – velké", "cena_mo": 299.9, "foto": "", "kategorie": "💖 Láska & Zvířátka"},
    {"nazev": "❤️ Srdíčko z lásky", "cena_mo": 249.9, "foto": "", "kategorie": "💖 Láska & Zvířátka"},
    {"nazev": "🧸 Medvídek z lásky", "cena_mo": 299.9, "foto": "", "kategorie": "💖 Láska & Zvířátka"},
    {"nazev": "🫂 Láskyplné objetí", "cena_mo": 299.9, "foto": "", "kategorie": "💖 Láska & Zvířátka"},
    {"nazev": "🐰 Velikonoční ušák – fialový", "cena_mo": 299.0, "foto": "", "kategorie": "💖 Láska & Zvířátka"}
]

st.set_page_config(page_title="Deni Candle | B2B Velkoobchod", layout="wide", page_icon="🕯️")

st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    [data-testid="stHeaderActionElements"] {visibility: hidden;}
    header {background-color: transparent !important;}
    footer {visibility: hidden;}
    
    /* Celkové pozadí - jemná pudrová růžová z webu */
    .stApp, [data-testid="stAppViewContainer"], [data-testid="stSidebar"] {
        background-color: #EBA299 !important;
    }
    
    html, body, [data-testid="stAppViewContainer"] *, [data-testid="stSidebar"] * {
        color: #261718 !important;
    }
    
    h1, h2, h3, h4 {
        color: #1F1012 !important;
        font-family: 'Playfair Display', 'Georgia', serif;
        font-weight: 700;
    }

    /* Čisté bílé karty pro katalog a košík */
    div[data-testid="stColumn"] {
        background: #FFFFFF !important;
        padding: 24px;
        border-radius: 16px;
        border: none !important;
        box-shadow: 0 8px 24px rgba(60, 20, 25, 0.08);
    }
    
    /* Vstupní pole */
    div[data-testid="stNumberInput"] div[data-baseweb="input"],
    div[data-baseweb="input"], input, textarea {
        background-color: #FFF9F8 !important;
        color: #261718 !important;
        -webkit-text-fill-color: #261718 !important;
        border: 1px solid #E5C2BD !important;
        border-radius: 10px !important;
    }
    
    div[data-baseweb="input"] button,
    div[data-testid="stNumberInput"] button {
        background-color: #F8E2DF !important;
        border: none !important;
        color: #261718 !important;
    }
    div[data-testid="stNumberInput"] button:hover {
        background-color: #EFC5C0 !important;
    }

    /* Tlačítka - signifikantní pastelově žlutá z webu */
    div.stButton > button,
    div.stDownloadButton > button,
    div.stFormSubmitButton > button,
    [data-testid="stDownloadButton"] > button,
    [data-testid="stFormSubmitButton"] > button {
        background-color: #FEE066 !important;
        color: #1A1A1A !important;
        border: 1px solid #EDCE4E !important;
        border-radius: 30px !important;
        font-weight: 700 !important;
        padding: 10px 24px !important;
        box-shadow: 0 4px 14px rgba(0,0,0,0.06);
        transition: all 0.2s ease-in-out;
    }
    div.stButton > button *,
    div.stDownloadButton > button *,
    div.stFormSubmitButton > button *,
    [data-testid="stDownloadButton"] > button *,
    [data-testid="stFormSubmitButton"] > button * {
        color: #1A1A1A !important;
    }
    div.stButton > button:hover,
    div.stDownloadButton > button:hover,
    div.stFormSubmitButton > button:hover {
        background-color: #FAD43D !important;
        transform: translateY(-1px);
        box-shadow: 0 6px 18px rgba(0,0,0,0.1);
    }

    /* Karty souhrnu a info boxy */
    .summary-card {
        background-color: #FFF2F0 !important;
        padding: 18px;
        border-radius: 12px;
        border-left: 6px solid #D87B70;
        margin-bottom: 20px;
    }

    /* Záložky (Tabs) */
    button[data-baseweb="tab"] {
        background-color: transparent !important;
        font-weight: 600 !important;
        color: #4A2E32 !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] {
        color: #1F1012 !important;
        border-bottom-color: #1F1012 !important;
    }

    /* Záhlaví ve stylu webu */
    .web-header-banner {
        background: linear-gradient(180deg, #100E0E 0%, #1A1616 100%);
        border-radius: 14px;
        padding: 24px;
        text-align: center;
        color: #FFFFFF !important;
        margin-bottom: 25px;
        box-shadow: 0 6px 20px rgba(0,0,0,0.25);
    }
    .web-header-banner * {
        color: #FFFFFF !important;
    }
    .web-header-contacts {
        font-size: 13px;
        letter-spacing: 1px;
        opacity: 0.85;
        margin-bottom: 12px;
    }
    </style>
""", unsafe_allow_html=True)

# Vizuální záhlaví ladící s webem
st.markdown("""
    <div class="web-header-banner">
        <div class="web-header-contacts">
            ✉️ kouba20@centrum.cz &nbsp;|&nbsp; 📞 +420 725 677 701 &nbsp;|&nbsp; 🇨🇿 Česká ruční výroba
        </div>
        <h1 style="margin: 0; font-size: 32px; letter-spacing: 2px;">🕯️ DENI CANDLE</h1>
        <p style="margin: 5px 0 0 0; font-size: 15px; opacity: 0.9;">Sojové svíčky s čistým svědomím & B2B velkoobchodní portál</p>
    </div>
""", unsafe_allow_html=True)

def get_headers():
    return {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}

def odeslat_email_upozorneni(id_obj, firma, jmeno, email_partnera, telefon, adresa, cena, doprava_nazev, celkem_ks, polozky_text, datum_dodani):
    if not EMAIL_SENDER or not EMAIL_PASSWORD:
        return False
    try:
        msg = MIMEMultipart()
        msg['From'] = f"Deni Candle <{EMAIL_SENDER}>"
        
        if email_partnera:
            msg['To'] = email_partnera
            if EMAIL_RECEIVER and EMAIL_RECEIVER.strip().lower() != email_partnera.strip().lower():
                msg['Cc'] = EMAIL_RECEIVER
        else:
            msg['To'] = EMAIL_RECEIVER

        msg['Subject'] = f"🕯️ Potvrzení B2B objednávky #{id_obj} — Deni Candle"
        
        body = f"""Dobrý den,

děkujeme za Vaši velkoobchodní objednávku v portálu Deni Candle.

📋 SHRNUTÍ OBJEDNÁVKY #{id_obj}
----------------------------------------
Firma / IČO: {firma}
Kontaktní osoba: {jmeno} ({telefon})
Doručovací adresa: {adresa}
Požadovaný termín dodání: {datum_dodani}
Způsob dopravy: {doprava_nazev} (bude vyčíslena dle celkového objemu balíku)
Celkem produktů: {celkem_ks} ks

PŘEHLED POLOŽEK:
{polozky_text}

----------------------------------------
CELKOVÁ CENA ZBOŽÍ: {cena:,.0f} Kč (+ doprava dle množství)
----------------------------------------

💳 POKYNY K PLATBĚ PŘEVODEM:
Číslo účtu: {BANK_ACCOUNT}/{BANK_CODE}
IBAN: {BANK_IBAN}
Variabilní symbol: {id_obj}
Částka za zboží: {cena:,.0f} Kč
Zpráva pro příjemce: DeniCandle B2B {id_obj}

Jakmile platbu přijmeme, pustíme se do odlévání a přípravy Vašeho balíku.

V případě jakýchkoli dotazů se na nás neváhejte obrátit.

S přáním krásného dne,
Deni Candle
"""
        msg.attach(MIMEText(body, 'plain', 'utf-8'))
        
        if SMTP_PORT == 465:
            server = smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT)
        else:
            server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
            server.starttls()
            
        server.login(EMAIL_SENDER, EMAIL_PASSWORD)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        print(f"Chyba pri odesilani e-mailu: {e}")
        return False

def nacti_nastaveni():
    vychozi_data = {
        "partneri": [
            {
                "PIN": VÝCHOZÍ_B2B_PIN, 
                "Nazev": "Základní velkoodběratel", 
                "Firma_ICO": "Svíčky s.r.o., IČO: 12345678",
                "Jmeno": "Jana Nováková",
                "Telefon": "+420 777 123 456",
                "Email": "info@partner.cz",
                "Adresa": "Nádražní 15, 301 00 Plzeň",
                "Sleva": 40
            }
        ],
        "vyrobeno_sklad": {}
    }
    if not GITHUB_TOKEN or not GITHUB_REPO:
        if not os.path.exists(SETTINGS_PATH):
            with open(SETTINGS_PATH, "w") as f:
                json.dump(vychozi_data, f)
        with open(SETTINGS_PATH, "r") as f:
            data = json.load(f)
            if "vyrobeno_sklad" not in data:
                data["vyrobeno_sklad"] = {}
            return data, None

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{SETTINGS_PATH}"
    res = requests.get(url, headers=get_headers())
    if res.status_code == 200:
        data = res.json()
        sha = data["sha"]
        content_str = base64.b64decode(data["content"]).decode("utf-8")
        nastaveni_json = json.loads(content_str)
        if "vyrobeno_sklad" not in nastaveni_json:
            nastaveni_json["vyrobeno_sklad"] = {}
        return nastaveni_json, sha
    else:
        return vychozi_data, None

def uloz_nastaveni(nastaveni_dict, sha=None):
    obsah = json.dumps(nastaveni_dict)
    if not GITHUB_TOKEN or not GITHUB_REPO:
        with open(SETTINGS_PATH, "w") as f:
            f.write(obsah)
        return True

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{SETTINGS_PATH}"
    content_b64 = base64.b64encode(obsah.encode("utf-8")).decode("utf-8")
    payload = {"message": "Aktualizace nastaveni / skladu Deni Candle", "content": content_b64}
    if sha:
        payload["sha"] = sha
    res = requests.put(url, headers=get_headers(), json=payload)
    return res.status_code in [200, 201]

def nacti_objednavky():
    sloupce = [
        "ID", "Datum_Vytvoreni", "Pozadovane_Datum_Dodani", "Oznaceni_Partnera", "Firma_ICO", "Jmeno_Kontakt", 
        "Telefon", "Email", "Adresa_Doruceni", "Poznamka", "Polozky_Detail", 
        "Celkem_Ks", "Sleva_Pouzita", "Cena_Zbozi_VO", "Doprava_Nazev", "Cena_Dopravy", 
        "Cena_Celkem_VO", "Stav_Platby", "Stav_Vyroby"
    ]
    
    if not GITHUB_TOKEN or not GITHUB_REPO:
        if not os.path.exists(FILE_PATH):
            df_empty = pd.DataFrame(columns=sloupce)
            df_empty.to_csv(FILE_PATH, index=False)
            return df_empty, None
        df = pd.read_csv(FILE_PATH)
        for col in ["Stav_Vyroby", "Doprava_Nazev", "Cena_Zbozi_VO", "Pozadovane_Datum_Dodani"]:
            if col not in df.columns:
                if col == "Pozadovane_Datum_Dodani":
                    df[col] = "-"
                elif col == "Stav_Vyroby":
                    df[col] = "K výrobě"
                elif col == "Doprava_Nazev":
                    df[col] = "Dle množství položek"
                else:
                    df[col] = df.get("Cena_Celkem_VO", 0)
        return df, None

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
    res = requests.get(url, headers=get_headers())
    if res.status_code == 200:
        data = res.json()
        sha = data["sha"]
        content_str = base64.b64decode(data["content"]).decode("utf-8")
        df = pd.read_csv(io.StringIO(content_str))
        for col in ["Stav_Vyroby", "Doprava_Nazev", "Cena_Zbozi_VO", "Pozadovane_Datum_Dodani"]:
            if col not in df.columns:
                if col == "Pozadovane_Datum_Dodani":
                    df[col] = "-"
                elif col == "Stav_Vyroby":
                    df[col] = "K výrobě"
                elif col == "Doprava_Nazev":
                    df[col] = "Dle množství položek"
                else:
                    df[col] = df.get("Cena_Celkem_VO", 0)
        return df, sha
    else:
        df_empty = pd.DataFrame(columns=sloupce)
        return df_empty, None

def uloz_objednavky(df, sha=None):
    if not GITHUB_TOKEN or not GITHUB_REPO:
        df.to_csv(FILE_PATH, index=False)
        return True

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
    csv_buffer = io.StringIO()
    df.to_csv(csv_buffer, index=False)
    content_b64 = base64.b64encode(csv_buffer.getvalue().encode("utf-8")).decode("utf-8")

    payload = {"message": "Aktualizace B2B objednavek Deni Candle", "content": content_b64}
    if sha:
        payload["sha"] = sha
    res = requests.put(url, headers=get_headers(), json=payload)
    return res.status_code in [200, 201]

def vytvor_profi_excel(df, titulek="Objednávky"):
    wb = Workbook()
    ws = wb.active
    ws.title = titulek
    
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
        
    hlavicka_fill = PatternFill(start_color="2A1B1D", end_color="2A1B1D", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    for cell in ws[1]:
        cell.fill = hlavicka_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col_letter].width = (max_length + 2)
        
    ws.auto_filter.ref = ws.dimensions
    
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def vygeneruj_b2b_uctenku(id_obj, firma, jmeno, adresa, telefon, polozky_str, celkem_ks, cena_zbozi, doprava_nazev, cena_celkem, sleva, datum_dodani):
    return f"""
    <!DOCTYPE html>
    <html lang="cs">
    <head>
    <meta charset="UTF-8">
    <title>B2B Objednávka Deni Candle #{id_obj}</title>
    <style>
        body {{ font-family: Arial, sans-serif; color: #111; max-width: 650px; margin: 0 auto; padding: 30px; background: #EBA299; }}
        .box {{ background: #fff; padding: 30px; border-radius: 12px; border-top: 8px solid #1A1A1A; box-shadow: 0 4px 15px rgba(0,0,0,0.1); }}
        .header {{ text-align: center; border-bottom: 2px solid #eee; padding-bottom: 15px; }}
        .header h1 {{ color: #1A1A1A; margin: 0; }}
        .total {{ background: #FFF2F0; padding: 15px; text-align: center; font-size: 20px; font-weight: bold; color: #2A1B1D; margin-top: 20px; border-radius: 6px; }}
        .payment {{ background: #FFFDF0; padding: 15px; border-radius: 6px; border-left: 4px solid #FEE066; margin-top: 20px; font-size: 14px; }}
    </style>
    </head>
    <body>
        <div class="box">
            <div class="header">
                <h1>🕯️ Deni Candle — B2B Potvrzení</h1>
                <p>Velkoobchodní objednávka #{id_obj} (Sleva {sleva} %)</p>
            </div>
            <h3>Odběratel:</h3>
            <p><strong>Firma / IČO:</strong> {firma}<br><strong>Kontakt:</strong> {jmeno} ({telefon})<br><strong>Adresa doručení:</strong> {adresa}<br><strong>Požadovaný termín dodání:</strong> {datum_dodani}</p>
            <h3>Objednané zboží ({celkem_ks} ks):</h3>
            <p>{polozky_str}</p>
            <p><strong>Doprava:</strong> {doprava_nazev} (bude vyčíslena dle celkového objemu)</p>
            <div class="total">Cena zboží k úhradě: {cena_celkem:,.0f} Kč</div>
            
            <div class="payment">
                <strong>💳 Pokyny k platbě převodem:</strong><br><br>
                Číslo účtu: <strong>{BANK_ACCOUNT}/{BANK_CODE}</strong><br>
                IBAN: <strong>{BANK_IBAN}</strong><br>
                Variabilní symbol: <strong>{id_obj}</strong><br>
                Částka za zboží: <strong>{cena_celkem:,.0f} Kč</strong>
            </div>
            
            <p style="margin-top:20px; font-size:12px; color:#666; text-align:center;">Děkujeme za váš odběr rukodělných svíček Deni Candle.</p>
        </div>
    </body>
    </html>
    """

df_orders, current_sha = nacti_objednavky()
nastaveni_app, sha_nastaveni = nacti_nastaveni()

partneri_seznam = nastaveni_app.get("partneri", [{"PIN": VÝCHOZÍ_B2B_PIN, "Nazev": "Základní velkoodběratel", "Sleva": 40}])
partneri_dict = {str(p["PIN"]): p for p in partneri_seznam}
sklad_vyrobeno = nastaveni_app.get("vyrobeno_sklad", {})

if os.path.exists(LOGO_PATH):
    st.sidebar.image(LOGO_PATH, width=140)

st.sidebar.markdown("## 🕯️ Deni Candle B2B")
rezim = st.sidebar.radio("Navigace:", ["🛍️ Velkoobchodní objednávka", "🔐 Správa pro majitele"])

# ---------------------------------------------------------
# 1. B2B OBJEDNÁVKOVÝ PORTÁL PRO PARTNERY
# ---------------------------------------------------------
if rezim == "🛍️ Velkoobchodní objednávka":
    if st.session_state.get("need_clear_b2b_form", False):
        for idx in range(len(PRODUKTY_KATALOG)):
            st.session_state[f"vo_{idx}"] = 0
        st.session_state["need_clear_b2b_form"] = False
    
    if "potvrzeni_b2b" in st.session_state:
        p = st.session_state["potvrzeni_b2b"]
        st.balloons()
        st.success(f"🎉 Objednávka #{p['id']} byla úspěšně odeslána! Potvrzení s platebními údaji bylo odesláno na Váš e-mail.")
        
        col_c1, col_c2 = st.columns([1.2, 1])
        with col_c1:
            st.download_button(
                "📥 Stáhnout B2B Potvrzení objednávky", 
                p["html_uct"], 
                file_name=f"DeniCandle_B2B_{p['id']}.html", 
                mime="text/html",
                key="btn_download_b2b_ok"
            )
            st.write(f"**Číslo účtu:** {BANK_ACCOUNT}/{BANK_CODE}")
            st.write(f"**IBAN:** {BANK_IBAN}")
            st.write(f"**Částka k úhradě:** {p['cena']:,.0f} Kč")
            st.write(f"**Variabilní symbol:** {p['id']}")
            st.write(f"**Zpráva pro příjemce:** DeniCandle B2B {p['id']}")
            
            if st.button("❌ Zavřít toto potvrzení", key="btn_close_b2b_confirmation"):
                del st.session_state["potvrzeni_b2b"]
                st.rerun()
                
        with col_c2:
            st.image(p["qr_bytes"], caption="Naskenujte v bankovní aplikaci", width=180)
            
        st.divider()
    
    url_pin = st.query_params.get("pin", "")
    zadan_pin = st.text_input("Zadejte přístupový B2B PIN partnera:", type="password", value=url_pin)
    
    if zadan_pin in partneri_dict:
        aktivni_partner = partneri_dict[zadan_pin]
        aktualni_sleva = aktivni_partner["Sleva"]
        nazev_partnera = aktivni_partner.get("Nazev", "")
        
        st.success(f"✅ Vítejte, **{nazev_partnera}**! Uplatněna vaše partnerská sleva **{aktualni_sleva} %** z maloobchodních cen.")
        
        tab_p1, tab_p2 = st.tabs(["🛒 Nová objednávka", "📜 Moje historie objednávek"])
        
        with tab_p1:
            col_katalog, col_kosik = st.columns([1.3, 1])
            
            vybrane_polozky = []
            celkova_cena_zbozi_vo = 0
            celkem_ks = 0
            mo_hodnota_celkem = 0
            
            with col_katalog:
                st.subheader("1. Výběr svíček podle kolekcí")
                
                kategorie_seznam = [
                    "🎄 Vánoční kolekce",
                    "🎃 Podzim & Halloween",
                    "🌸 Květinové & Elegantní",
                    "🕯️ Dárkové & Motivační",
                    "💖 Láska & Zvířátka"
                ]
                
                tabs_kategorie = st.tabs(kategorie_seznam)
                for tab_k, kat_nazev in zip(tabs_kategorie, kategorie_seznam):
                    with tab_k:
                        polozky_v_kat = [(idx, prod) for idx, prod in enumerate(PRODUKTY_KATALOG) if prod["kategorie"] == kat_nazev]
                        col_p1, col_p2 = st.columns(2)
                        
                        for i, (idx, prod) in enumerate(polozky_v_kat):
                            target_col = col_p1 if i % 2 == 0 else col_p2
                            cena_mo = prod["cena_mo"]
                            cena_vo = round(cena_mo * (1 - aktualni_sleva / 100))
                            
                            if prod.get("foto") and os.path.exists(prod["foto"]):
                                target_col.image(prod["foto"], width=100)
                            
                            label = f"{prod['nazev']} (VO: {cena_vo} Kč | MO: {cena_mo:.0f} Kč)"
                            key_input = f"vo_{idx}"
                            default_val = st.session_state.get(key_input, 0)
                            ks = target_col.number_input(label, min_value=0, max_value=500, value=default_val, key=key_input)
                            
                            if ks > 0:
                                vybrane_polozky.append(f"{ks}x {prod['nazev']}")
                                celkova_cena_zbozi_vo += ks * cena_vo
                                mo_hodnota_celkem += ks * cena_mo
                                celkem_ks += ks

            with col_kosik:
                st.subheader("2. Údaje & Dodání")
                
                vychozi_firma = aktivni_partner.get("Firma_ICO") or aktivni_partner.get("Nazev", "")
                vychozi_jmeno = aktivni_partner.get("Jmeno", "")
                vychozi_telefon = aktivni_partner.get("Telefon", "")
                vychozi_email = aktivni_partner.get("Email", "")
                vychozi_adresa = aktivni_partner.get("Adresa", "")
                
                firma = st.text_input("Název firmy / Obchodu & IČO *", value=vychozi_firma).strip()
                jmeno = st.text_input("Jméno kontaktní osoby *", value=vychozi_jmeno).strip()
                col_t, col_e = st.columns(2)
                telefon = col_t.text_input("Telefon *", value=vychozi_telefon).strip()
                email = col_e.text_input("E-mail * (pro potvrzení objednávky)", value=vychozi_email).strip()
                adresa = st.text_area("Doručovací adresa (Ulice, ČP, Město, PSČ) *", value=vychozi_adresa, height=80).strip()
                
                min_datum_dodani = date.today() + timedelta(days=4)
                datum_dodani = st.date_input(
                    "📅 Požadovaný termín dodání * (výroba zabere min. 4 dny):",
                    value=min_datum_dodani,
                    min_value=min_datum_dodani,
                    format="DD.MM.YYYY"
                )
                datum_dodani_str = datum_dodani.strftime("%d.%m.%Y")
                
                st.markdown("**🚚 Způsob a cena dopravy:**")
                st.info("📦 **Doprava bude vybrána dle množství položek a cena bude také stanovena dle celkového objemu/hmotnosti zásilky.** (Bude dohodnuto při expedici).")
                zvolena_doprava_nazev = "Dle množství položek"
                cena_dopravy = 0

                poznamka = st.text_input("Poznámka k doručení / balení:").strip()
                
                st.divider()
                
                celkova_konecna_cena = celkova_cena_zbozi_vo
                uspora = mo_hodnota_celkem - celkova_cena_zbozi_vo
                
                st.markdown(f"""
                    <div class='summary-card'>
                        <h4 style='margin: 0;'>Shrnutí VO objednávky</h4>
                        <p style='margin: 5px 0 0 0; font-size: 14px;'>Celkem produktů: <strong>{celkem_ks} ks</strong></p>
                        <p style='margin: 0; font-size: 13px;'>Požadované dodání: <strong>{datum_dodani_str}</strong></p>
                        <p style='margin: 0; font-size: 13px;'>Doprava: <em>Dle množství položek (bude dohodnuta)</em></p>
                        <p style='margin: 0; font-size: 13px; color: #8C5A47;'>Úspora oproti MO: {uspora:,.0f} Kč</p>
                        <h3 style='margin: 10px 0 0 0;'>Cena zboží k úhradě: {celkova_konecna_cena:,.0f} Kč</h3>
                        <small style='color: #666;'>+ cena dopravy dle množství položek</small>
                    </div>
                """, unsafe_allow_html=True)
                
                if st.button("Odeslat velkoobchodní objednávku", type="primary", use_container_width=True):
                    if not all([firma, jmeno, telefon, email, adresa]):
                        st.warning("⚠️ Prosím vyplňte všechny kontaktní a firemní údaje.")
                    elif celkem_ks == 0:
                        st.error("❌ Košík je prázdný. Vyberte prosím alespoň jednu svíčku.")
                    else:
                        with st.spinner('Odesílám VO objednávku... 🕯️'):
                            nove_id = 1 if df_orders.empty else int(df_orders["ID"].max()) + 1
                            polozky_text_email = ",\n".join(vybrane_polozky)
                            polozky_text_db = ", ".join(vybrane_polozky)
                            
                            nova_obj = pd.DataFrame([{
                                "ID": nove_id,
                                "Datum_Vytvoreni": datetime.now().strftime("%Y-%m-%d %H:%M"),
                                "Pozadovane_Datum_Dodani": datum_dodani_str,
                                "Oznaceni_Partnera": nazev_partnera,
                                "Firma_ICO": firma,
                                "Jmeno_Kontakt": jmeno,
                                "Telefon": telefon,
                                "Email": email,
                                "Adresa_Doruceni": adresa,
                                "Poznamka": poznamka,
                                "Polozky_Detail": polozky_text_db,
                                "Celkem_Ks": celkem_ks,
                                "Sleva_Pouzita": f"{aktualni_sleva} %",
                                "Cena_Zbozi_VO": celkova_cena_zbozi_vo,
                                "Doprava_Nazev": zvolena_doprava_nazev,
                                "Cena_Dopravy": cena_dopravy,
                                "Cena_Celkem_VO": celkova_konecna_cena,
                                "Stav_Platby": "Čeká na platbu",
                                "Stav_Vyroby": "K výrobě"
                            }])
                            
                            df_aktualni = pd.concat([df_orders, nova_obj], ignore_index=True)
                            if uloz_objednavky(df_aktualni, current_sha):
                                odeslat_email_upozorneni(
                                    id_obj=nove_id, 
                                    firma=firma, 
                                    jmeno=jmeno, 
                                    email_partnera=email, 
                                    telefon=telefon, 
                                    adresa=adresa, 
                                    cena=celkova_konecna_cena, 
                                    doprava_nazev=zvolena_doprava_nazev, 
                                    celkem_ks=celkem_ks, 
                                    polozky_text=polozky_text_email,
                                    datum_dodani=datum_dodani_str
                                )
                                
                                html_uct = vygeneruj_b2b_uctenku(
                                    nove_id, firma, jmeno, adresa, telefon, 
                                    polozky_text_email.replace("\n", "<br>"), 
                                    celkem_ks, celkova_cena_zbozi_vo, zvolena_doprava_nazev, 
                                    celkova_konecna_cena, aktualni_sleva, datum_dodani_str
                                )
                                
                                spd_str = f"SPD*1.0*ACC:{BANK_IBAN}*AM:{celkova_konecna_cena:.2f}*CC:CZK*X-VS:{nove_id}*MSG:DeniCandle B2B {nove_id}"
                                qr_img = qrcode.make(spd_str)
                                buf = BytesIO()
                                qr_img.save(buf, format="PNG")
                                
                                st.session_state["potvrzeni_b2b"] = {
                                    "id": nove_id,
                                    "html_uct": html_uct,
                                    "qr_bytes": buf.getvalue(),
                                    "cena": celkova_konecna_cena
                                }
                                st.session_state["need_clear_b2b_form"] = True
                                st.rerun()
                            else:
                                st.error("❌ Chyba při ukládání objednávky.")
                                
        with tab_p2:
            st.markdown(f"### 📜 Minulé objednávky partnera **{nazev_partnera}**")
            moje_objednavky = df_orders[df_orders['Oznaceni_Partnera'] == nazev_partnera].sort_values(by="ID", ascending=False) if not df_orders.empty else pd.DataFrame()
            
            if not moje_objednavky.empty:
                for idx, row in moje_objednavky.iterrows():
                    id_o = row['ID']
                    dat_o = row['Datum_Vytvoreni']
                    dat_dod = row.get('Pozadovane_Datum_Dodani', '-')
                    cena_o = row['Cena_Celkem_VO']
                    ks_o = row['Celkem_Ks']
                    detail_o = str(row['Polozky_Detail'])
                    
                    with st.expander(f"📦 Objednávka #{id_o} — {cena_o:,.0f} Kč (Vytvořeno: {dat_o} | Dodání: {dat_dod})"):
                        st.write(f"**Položky ({ks_o} ks):** {detail_o}")
                        st.write(f"**Doprava:** {row.get('Doprava_Nazev', '-')} | **Stav:** {row.get('Stav_Vyroby', 'K výrobě')}")
                        
                        if st.button(f"🔄 Zopakovat objednávku #{id_o}", key=f"repeat_{id_o}"):
                            for i_p, prod in enumerate(PRODUKTY_KATALOG):
                                st.session_state[f"vo_{i_p}"] = 0
                                for item in detail_o.split(", "):
                                    if "x " in item:
                                        k_str, n_str = item.split("x ", 1)
                                        if n_str.strip() == prod["nazev"].strip():
                                            st.session_state[f"vo_{i_p}"] = int(k_str)
                            st.success(f"Položky z objednávky #{id_o} byly načteny do košíku! Přepněte se na záložku 'Nová objednávka'.")
                            st.rerun()
            else:
                st.info("Zatím jste nevytvořili žádné objednávky.")

    elif zadan_pin:
        st.error("❌ Neplatný B2B PIN. Zkontrolujte prosím přístupové heslo.")

# ---------------------------------------------------------
# 2. SPRÁVA PRO MAJITELE DENI CANDLE
# ---------------------------------------------------------
else:
    st.markdown("<h2>🔐 Správa Deni Candle</h2>", unsafe_allow_html=True)
    heslo_admin = st.text_input("Heslo správce:", type="password")
    
    if heslo_admin == ADMIN_PASSWORD:
        st.success("✅ Přístup schválen.")
        
        tab_db, tab_vyroba, tab_partneri = st.tabs(["📋 Databáze objednávek", "🕯️ Souhrn pro výrobu", "👥 Správa odběratelů a slev"])
        
        with tab_db:
            if not df_orders.empty:
                st.markdown("### 📋 Přehled přijatých objednávek")
                
                excel_db = vytvor_profi_excel(df_orders, titulek="Databaze_Objednavek")
                st.download_button(
                    label="📥 Stáhnout kompletní databázi do Excelu", 
                    data=excel_db, 
                    file_name=f"DeniCandle_Databaze_{datetime.now().strftime('%d_%m_%Y')}.xlsx", 
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_db_all"
                )
                
                st.divider()
                
                df_orders_sorted = df_orders.sort_values(by="ID", ascending=False)
                
                for idx, row in df_orders_sorted.iterrows():
                    id_obj = row['ID']
                    datum = row['Datum_Vytvoreni']
                    dodani = row.get('Pozadovane_Datum_Dodani', '-')
                    partner = row['Oznaceni_Partnera']
                    cena = row['Cena_Celkem_VO']
                    stav_v = row.get('Stav_Vyroby', 'K výrobě')
                    
                    ikona_stavu = "🔥" if stav_v == "K výrobě" else "✅"
                    
                    with st.expander(f"{ikona_stavu} Objednávka #{id_obj} — {partner} | {cena:,.0f} Kč | Dodání do: {dodani} ({stav_v})"):
                        c1, c2 = st.columns(2)
                        
                        with c1:
                            st.markdown(f"**🏢 Odběratel:** {row['Firma_ICO']}")
                            st.markdown(f"**👤 Kontakt:** {row['Jmeno_Kontakt']}")
                            st.markdown(f"**📞 Telefon:** {row['Telefon']}")
                            st.markdown(f"**✉️ E-mail:** {row['Email']}")
                            st.markdown(f"**📍 Adresa:** {row['Adresa_Doruceni']}")
                            
                        with c2:
                            st.markdown(f"**📅 Požadované datum dodání:** `{dodani}`")
                            st.markdown(f"**🏷️ Použitá sleva:** {row['Sleva_Pouzita']}")
                            st.markdown(f"**🚚 Doprava:** {row.get('Doprava_Nazev', 'Dle množství položek')}")
                            st.markdown(f"**💰 Cena zboží (VO):** {cena:,.0f} Kč")
                            st.markdown(f"**📦 Celkem kusů:** {row['Celkem_Ks']} ks")
                            st.markdown(f"**📌 Stav zakázky:** `{stav_v}`")
                            
                            poznamka = row['Poznamka']
                            if pd.notna(poznamka) and str(poznamka).strip() != "":
                                st.markdown(f"**📝 Poznámka:** {poznamka}")
                        
                        st.markdown("**🛍️ Objednané položky:**")
                        st.info(row['Polozky_Detail'])
                        
                        col_b1, col_b2 = st.columns(2)
                        if stav_v == "K výrobě":
                            if col_b1.button(f"✅ Označit #{id_obj} jako vyřízenou/odlitou", key=f"done_db_{id_obj}"):
                                df_orders.loc[df_orders['ID'] == id_obj, 'Stav_Vyroby'] = 'Vyřízeno'
                                uloz_objednavky(df_orders, current_sha)
                                st.success("✅ Změněno na Vyřízeno.")
                                st.rerun()
                        else:
                            if col_b1.button(f"🔄 Vrátit #{id_obj} zpět k výrobě", key=f"reopen_db_{id_obj}"):
                                df_orders.loc[df_orders['ID'] == id_obj, 'Stav_Vyroby'] = 'K výrobě'
                                uloz_objednavky(df_orders, current_sha)
                                st.rerun()
                                
                        if col_b2.button(f"🗑️ Smazat objednávku #{id_obj}", key=f"del_db_{id_obj}"):
                            df_orders = df_orders[df_orders['ID'] != id_obj]
                            uloz_objednavky(df_orders, current_sha)
                            st.success("🗑️ Objednávka smazána.")
                            st.rerun()
            else:
                st.info("Zatím žádné B2B objednávky.")
                
        with tab_vyroba:
            df_k_vyrobe = df_orders[df_orders.get('Stav_Vyroby', 'K výrobě') == 'K výrobě'] if not df_orders.empty else pd.DataFrame()
            
            if not df_k_vyrobe.empty:
                st.markdown("### 📊 Interaktivní výrobní matice (Přehled chybějících kusů)")
                st.info("💡 **Nápověda:** V tabulce níže dvakrát klikněte do sloupce **'Už vyrobeno'** a zadejte aktuální počet odlitých svíček. Systém sám spočítá, kolik jich ještě chybí pro pokrytí všech otevřených objednávek.")
                
                matrix_rows = []
                for idx, row in df_k_vyrobe.iterrows():
                    id_obj = row['ID']
                    partner = row['Oznaceni_Partnera']
                    col_label = f"Obj #{id_obj} ({partner})"
                    
                    detail = row["Polozky_Detail"]
                    if pd.notna(detail):
                        for item in str(detail).split(", "):
                            if "x " in item:
                                ks, název = item.split("x ", 1)
                                matrix_rows.append({"Produkt": název, "Col": col_label, "Ks": int(ks)})
                
                if matrix_rows:
                    df_matrix_raw = pd.DataFrame(matrix_rows)
                    df_pivot = df_matrix_raw.pivot_table(index="Produkt", columns="Col", values="Ks", aggfunc="sum", fill_value=0)
                    
                    df_pivot["CELKEM POTŘEBA"] = df_pivot.sum(axis=1)
                    df_pivot["Už vyrobeno"] = df_pivot.index.map(lambda x: sklad_vyrobeno.get(x, 0))
                    df_pivot["Zbývá odlejt"] = df_pivot["CELKEM POTŘEBA"] - df_pivot["Už vyrobeno"]
                    df_pivot["Zbývá odlejt"] = df_pivot["Zbývá odlejt"].apply(lambda x: x if x > 0 else 0)
                    
                    cols = df_pivot.columns.tolist()
                    order_cols = [c for c in cols if c.startswith("Obj #")]
                    new_cols = ["CELKEM POTŘEBA", "Už vyrobeno", "Zbývá odlejt"] + order_cols
                    df_pivot = df_pivot[new_cols].sort_values(by="Zbývá odlejt", ascending=False).reset_index()
                    
                    edited_df = st.data_editor(
                        df_pivot,
                        hide_index=True,
                        column_config={
                            "Produkt": st.column_config.TextColumn("Produkt", disabled=True),
                            "CELKEM POTŘEBA": st.column_config.NumberColumn("CELKEM POTŘEBA", disabled=True),
                            "Zbývá odlejt": st.column_config.NumberColumn("Zbývá odlejt (ks) 🔥", disabled=True),
                            "Už vyrobeno": st.column_config.NumberColumn(
                                "Už vyrobeno (ks) ✍️", 
                                help="Klikněte sem pro úpravu", 
                                min_value=0,
                                step=1
                            )
                        },
                        disabled=["Produkt", "CELKEM POTŘEBA", "Zbývá odlejt"] + order_cols,
                        key="editor_sklad",
                        use_container_width=True
                    )
                    
                    if st.button("💾 Uložit aktuální stav výroby", type="primary"):
                        for idx, row in edited_df.iterrows():
                            produkt_nazev = row["Produkt"]
                            vyrobeno_ks = int(row["Už vyrobeno"])
                            sklad_vyrobeno[produkt_nazev] = vyrobeno_ks
                        
                        nastaveni_app["vyrobeno_sklad"] = sklad_vyrobeno
                        if uloz_nastaveni(nastaveni_app, sha_nastaveni):
                            st.success("✅ Stav výroby byl úspěšně uložen.")
                            st.rerun()

                    st.divider()
                    
                    excel_data = vytvor_profi_excel(edited_df, titulek="Vyrobni_Matice")
                    st.download_button(
                        label="📥 Stáhnout tuto matici do Excelu", 
                        data=excel_data, 
                        file_name=f"DeniCandle_Vyroba_{datetime.now().strftime('%d_%m')}.xlsx", 
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_all_vyroba_matrix"
                    )
                else:
                    st.info("Žádné svíčky k výrobě.")
                
                st.divider()
                
                st.markdown("### 📦 Samostatné rozpisky k odlítí")
                df_k_vyrobe_sorted = df_k_vyrobe.sort_values(by="ID", ascending=False)
                
                for idx, row in df_k_vyrobe_sorted.iterrows():
                    id_obj = row['ID']
                    partner = row['Oznaceni_Partnera']
                    datum = row['Datum_Vytvoreni']
                    dodani = row.get('Pozadovane_Datum_Dodani', '-')
                    
                    with st.expander(f"Rozpiska pro Objednávku #{id_obj} — {partner} (Termín: {dodani})"):
                        polozky_obj = []
                        detail = row["Polozky_Detail"]
                        if pd.notna(detail):
                            for item in str(detail).split(", "):
                                if "x " in item:
                                    ks, název = item.split("x ", 1)
                                    polozky_obj.append({"Produkt": název, "Ks": int(ks)})
                        
                        if polozky_obj:
                            df_obj = pd.DataFrame(polozky_obj).groupby("Produkt").sum().reset_index()
                            st.table(df_obj)
                            
                            c_m1, c_m2 = st.columns(2)
                            excel_obj = vytvor_profi_excel(df_obj, titulek=f"Objednavka_{id_obj}")
                            c_m1.download_button(
                                label=f"📥 Excel pro objednávku #{id_obj}", 
                                data=excel_obj, 
                                file_name=f"DeniCandle_Vyroba_Obj_{id_obj}.xlsx", 
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"dl_obj_{id_obj}"
                            )
                            
                            if c_m2.button(f"✅ Hotovo / Odlito (Vyřídit #{id_obj})", key=f"done_vyroba_{id_obj}"):
                                detail_odecet = row["Polozky_Detail"]
                                if pd.notna(detail_odecet):
                                    for item in str(detail_odecet).split(", "):
                                        if "x " in item:
                                            ks, n_str = item.split("x ", 1)
                                            if n_str in sklad_vyrobeno:
                                                novy_stav = sklad_vyrobeno[n_str] - int(ks)
                                                sklad_vyrobeno[n_str] = novy_stav if novy_stav > 0 else 0
                                    
                                    nastaveni_app["vyrobeno_sklad"] = sklad_vyrobeno
                                    uloz_nastaveni(nastaveni_app, sha_nastaveni)

                                df_orders.loc[df_orders['ID'] == id_obj, 'Stav_Vyroby'] = 'Vyřízeno'
                                uloz_objednavky(df_orders, current_sha)
                                st.success(f"Objednávka #{id_obj} vyřízena a sklad aktualizován.")
                                st.rerun()
            else:
                st.success("🎉 Skvělé! Všechny svíčky jsou odlité a žádné objednávky nečekají na výrobu.")
                
        with tab_partneri:
            st.markdown("### ➕ Založit nového velkoodběratele (Kompletní profil)")
            with st.form("form_novy_partner", clear_on_submit=True):
                col_p1, col_p2, col_p3 = st.columns([1.5, 2.5, 1])
                n_pin = col_p1.text_input("Přístupový PIN (heslo partnera) *").strip()
                n_nazev = col_p2.text_input("Zkrácený název partnera (např. Květiny Rosa) *").strip()
                n_sleva = col_p3.number_input("Sleva v %", min_value=0, max_value=99, value=40)
                
                col_f1, col_f2 = st.columns(2)
                n_firma = col_f1.text_input("Firma / Obchod & IČO * (pro fakturaci)").strip()
                n_jmeno = col_f2.text_input("Jméno kontaktní osoby *").strip()
                
                col_c1, col_c2 = st.columns(2)
                n_telefon = col_c1.text_input("Telefon *").strip()
                n_email = col_c2.text_input("E-mail *").strip()
                
                n_adresa = st.text_area("Výchozí doručovací adresa (Ulice, ČP, Město, PSČ) *", height=70).strip()
                
                btn_pridat = st.form_submit_button("💾 Uložit nového partnera")
                if btn_pridat:
                    if not n_pin or not n_nazev:
                        st.warning("⚠️ Vyplňte prosím alespoň PIN a název partnera.")
                    else:
                        existujici_piny = [str(p["PIN"]) for p in partneri_seznam]
                        if n_pin in existujici_piny:
                            st.error("❌ Tento PIN už používá jiný partner. Zvolte jiný.")
                        else:
                            novy_p = {
                                "PIN": n_pin,
                                "Nazev": n_nazev,
                                "Firma_ICO": n_firma if n_firma else n_nazev,
                                "Jmeno": n_jmeno,
                                "Telefon": n_telefon,
                                "Email": n_email,
                                "Adresa": n_adresa,
                                "Sleva": int(n_sleva)
                            }
                            partneri_seznam.append(novy_p)
                            if uloz_nastaveni({"partneri": partneri_seznam}, sha_nastaveni):
                                st.success(f"✅ Partner **{n_nazev}** byl úspěšně kompletně založen!")
                                st.rerun()
                            else:
                                st.error("❌ Chyba při ukládání.")

            st.divider()
            st.markdown("### 👥 Seznam aktivních velkoodběratelů")
            
            partneri_ke_smazani = []
            for i, p in enumerate(partneri_seznam):
                p_pin = p["PIN"]
                p_nazev = p.get("Nazev", "Partner")
                p_sleva = p.get("Sleva", 40)
                p_firma = p.get("Firma_ICO", "-")
                p_kontakt = p.get("Jmeno", "-")
                p_tel = p.get("Telefon", "-")
                p_mail = p.get("Email", "-")
                p_adr = p.get("Adresa", "-")
                
                with st.expander(f"🏪 {p_nazev} — Sleva {p_sleva} % (PIN: {p_pin})"):
                    c_info1, c_info2 = st.columns(2)
                    with c_info1:
                        st.markdown(f"**🔑 PIN:** `{p_pin}`")
                        st.markdown(f"**🏢 Firma / IČO:** {p_firma}")
                        st.markdown(f"**👤 Kontakt:** {p_kontakt}")
                    with c_info2:
                        st.markdown(f"**📞 Telefon:** {p_tel}")
                        st.markdown(f"**✉️ E-mail:** {p_mail}")
                        st.markdown(f"**📍 Adresa:** {p_adr}")
                        
                    st.caption(f"🔗 **Přímý odkaz pro partnera k odeslání:** `?pin={p_pin}`")
                    
                    if st.button("🗑️ Smazat tohoto partnera", key=f"del_partner_{i}"):
                        partneri_ke_smazani.append(i)
                    
            if partneri_ke_smazani:
                for idx in sorted(partneri_ke_smazani, reverse=True):
                    partneri_seznam.pop(idx)
                if uloz_nastaveni({"partneri": partneri_seznam}, sha_nastaveni):
                    st.success("✅ Partner byl smazán.")
                    st.rerun()
